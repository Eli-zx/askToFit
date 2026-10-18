"""多格式文档解析测试。"""

from pathlib import Path

import pytest
from PIL import Image

from app.services import parsers
from app.services.parsers import DocumentParseError, parse_document


def test_parse_markdown_with_sections(tmp_path: Path) -> None:
    path = tmp_path / "退换货规则.md"
    path.write_text(
        "# 退换货规则\n\n"
        "本规则适用于平台内所有服装商品。\n\n"
        "## 七天无理由\n\n"
        "自签收之日起七天内可申请无理由退换。\n\n"
        "## 质量问题\n\n"
        "出现质量问题时运费由商家承担。\n",
        encoding="utf-8",
    )
    parsed = parse_document(path, "md")
    assert parsed.title == "退换货规则"
    sections = [block.section for block in parsed.blocks]
    assert "七天无理由" in sections
    assert "质量问题" in sections
    assert any("七天内可申请无理由退换" in block.text for block in parsed.blocks)


def test_parse_markdown_empty_file_raises(tmp_path: Path) -> None:
    path = tmp_path / "empty.md"
    path.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(DocumentParseError):
        parse_document(path, "md")


def test_parse_pptx_extracts_slides_and_tables(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    path = tmp_path / "物流说明.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "物流时效"
    slide.placeholders[1].text = "江浙沪地区次日达。\n偏远地区三到五天。"
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "运费标准"
    table = table_slide.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    ).table
    table.cell(0, 0).text = "地区"
    table.cell(0, 1).text = "首重运费"
    table.cell(1, 0).text = "江浙沪"
    table.cell(1, 1).text = "8 元"
    presentation.save(path)

    parsed = parse_document(path, "pptx")
    assert parsed.title == "物流时效"
    assert len(parsed.blocks) == 2
    assert parsed.blocks[0].page == 1
    assert parsed.blocks[1].section == "运费标准"
    assert "江浙沪 | 8 元" in parsed.blocks[1].text


def test_parse_pdf_text_layer(tmp_path: Path) -> None:
    import fitz

    path = tmp_path / "invoice.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Invoice policy: e-invoice within 3 working days.")
    document.save(path)
    document.close()

    parsed = parse_document(path, "pdf")
    assert parsed.blocks
    assert any("e-invoice" in block.text for block in parsed.blocks)
    assert parsed.blocks[0].page == 1


def test_parse_pdf_scanned_page_uses_ocr(tmp_path: Path, monkeypatch) -> None:
    import fitz

    path = tmp_path / "scan.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()

    monkeypatch.setattr(parsers, "ocr_image", lambda image: ["扫描件识别出的退换货条款"])
    parsed = parse_document(path, "pdf")
    assert any("扫描件识别出的退换货条款" in block.text for block in parsed.blocks)


def test_parse_image_uses_ocr(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "policy.png"
    Image.new("RGB", (200, 80), "white").save(path)
    monkeypatch.setattr(
        parsers, "ocr_image", lambda image: ["支持货到付款", "满 199 元包邮"]
    )

    parsed = parse_document(path, "png")
    assert parsed.title == "policy"
    assert any("支持货到付款" in block.text for block in parsed.blocks)


def test_parse_image_prefers_title_hint(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "4b3db9f91f3548078a972723ef3462fd.png"
    Image.new("RGB", (200, 80), "white").save(path)
    monkeypatch.setattr(parsers, "ocr_image", lambda image: ["尺码速查", "女装 M 码"])

    parsed = parse_document(path, "png", "尺码速查")
    assert parsed.title == "尺码速查"
    assert all(block.title == "尺码速查" for block in parsed.blocks)


def test_parse_pdf_prefers_title_hint_when_metadata_missing(tmp_path: Path) -> None:
    import fitz

    path = tmp_path / "4b3db9f91f3548078a972723ef3462fd.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Refund arrives within three working days.")
    document.save(path)
    document.close()

    parsed = parse_document(path, "pdf", "退款到账说明")
    assert parsed.title == "退款到账说明"
    assert all(block.title == "退款到账说明" for block in parsed.blocks)


def test_parse_docx_extracts_headings_tables_and_paragraphs(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "退换货与售后政策手册.docx"
    document = DocxDocument()
    document.add_heading("退换货与售后政策手册", level=1)
    document.add_heading("七天无理由", level=2)
    document.add_paragraph("自签收之日起七天内，商品保持完好可申请无理由退换。")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "品类"
    table.cell(0, 1).text = "时限"
    table.cell(1, 0).text = "女装"
    table.cell(1, 1).text = "7 天"
    table.cell(2, 0).text = "定制西装"
    table.cell(2, 1).text = "不支持无理由"
    document.save(path)

    parsed = parse_document(path, "docx")
    assert parsed.title == "退换货与售后政策手册"
    assert any(block.section == "七天无理由" for block in parsed.blocks)
    assert any("自签收之日起七天内" in block.text for block in parsed.blocks)
    assert any("女装 | 7 天" in block.text for block in parsed.blocks)
    assert any("定制西装 | 不支持无理由" in block.text for block in parsed.blocks)


def test_parse_docx_uses_title_hint_without_heading(tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "uuid-name.docx"
    document = DocxDocument()
    document.add_paragraph("支持七天无理由退换，吊牌需保持完整。")
    document.save(path)

    parsed = parse_document(path, "docx", "退换货说明")
    assert parsed.title == "退换货说明"
    assert all(block.title == "退换货说明" for block in parsed.blocks)


def test_parse_docx_embedded_image_uses_ocr(tmp_path: Path, monkeypatch) -> None:
    from docx import Document as DocxDocument

    image_path = tmp_path / "size-chart.png"
    Image.new("RGB", (240, 120), "white").save(image_path)
    path = tmp_path / "尺码指南.docx"
    document = DocxDocument()
    document.add_heading("尺码指南", level=1)
    document.add_paragraph("尺码表以实际测量为准。")
    document.add_picture(str(image_path))
    document.save(path)

    monkeypatch.setattr(
        parsers,
        "ocr_image",
        lambda image: ["图片文字：", "女装 M 码 155-163cm", "女装 L 码 160-168cm"],
    )
    parsed = parse_document(path, "docx")
    assert any("女装 M 码 155-163cm" in block.text for block in parsed.blocks)


def test_parse_pptx_embedded_image_uses_ocr(tmp_path: Path, monkeypatch) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    image_path = tmp_path / "fabric.png"
    Image.new("RGB", (240, 120), "white").save(image_path)
    path = tmp_path / "新品面料.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "面料卖点"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(2), Inches(4), Inches(2))
    presentation.save(path)

    monkeypatch.setattr(
        parsers,
        "ocr_image",
        lambda image: ["桑蚕丝 19 姆米", "支持 30 度冷水手洗"],
    )
    parsed = parse_document(path, "pptx")
    assert parsed.title == "面料卖点"
    assert any("桑蚕丝 19 姆米" in block.text for block in parsed.blocks)


def test_parse_pdf_embedded_image_uses_ocr(tmp_path: Path, monkeypatch) -> None:
    import fitz

    image_path = tmp_path / "size-chart.png"
    Image.new("RGB", (240, 120), "white").save(image_path)
    path = tmp_path / "尺码与版型指南.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Size guide: measure your chest and waist before choosing.")
    page.insert_image(fitz.Rect(72, 120, 312, 240), filename=str(image_path))
    document.save(path)
    document.close()

    monkeypatch.setattr(
        parsers,
        "ocr_image",
        lambda image: ["女装 M 码：身高 155-163cm", "修身版型建议选大一码"],
    )
    parsed = parse_document(path, "pdf")
    assert any("女装 M 码" in block.text for block in parsed.blocks)


def test_parse_doc_uses_soffice_conversion(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy-word-placeholder")

    monkeypatch.setattr(
        parsers,
        "_docx_from_soffice",
        lambda file_path, title_hint: parsers.ParsedDocument(
            title="门店导购应答卡",
            blocks=[
                parsers.ParsedBlock(
                    text="顾客问：真丝可以机洗吗？答：建议冷水手洗。",
                    title="门店导购应答卡",
                )
            ],
        ),
    )
    parsed = parse_document(path, "doc")
    assert parsed.title == "门店导购应答卡"
    assert any("真丝" in block.text for block in parsed.blocks)


def test_parse_doc_falls_back_to_antiword(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"legacy-word-placeholder")

    monkeypatch.setattr(parsers, "_docx_from_soffice", lambda *args: None)
    monkeypatch.setattr(
        parsers,
        "_antiword_text",
        lambda file_path: "门店导购应答卡\n\n顾客问：支持七天无理由吗？\n答：吊牌完整即可。",
    )
    parsed = parse_document(path, "doc", "old")
    assert parsed.title == "old"
    assert any("吊牌完整即可" in block.text for block in parsed.blocks)


def test_parse_doc_raises_when_all_converters_fail(tmp_path: Path, monkeypatch) -> None:
    path = tmp_path / "broken.doc"
    path.write_bytes(b"not-a-word-file")

    monkeypatch.setattr(parsers, "_docx_from_soffice", lambda *args: None)
    monkeypatch.setattr(parsers, "_antiword_text", lambda file_path: "")
    monkeypatch.setattr(parsers, "_ole_doc_text", lambda file_path: "")
    with pytest.raises(DocumentParseError):
        parse_document(path, "doc")


def test_parse_unsupported_extension(tmp_path: Path) -> None:
    path = tmp_path / "data.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(DocumentParseError):
        parse_document(path, "txt")


def test_parse_missing_file(tmp_path: Path) -> None:
    with pytest.raises(DocumentParseError):
        parse_document(tmp_path / "missing.md", "md")
