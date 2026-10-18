"""多格式文档解析：Markdown / PDF / PPT(X) / Word / 图片。

解析结果按「块」返回，并携带页码、章节等元数据，供切块与检索引用。
重依赖（fitz、pptx、docx、rapidocr、cv2）均在函数内部懒加载，保证服务启动轻量。
"""

from __future__ import annotations

import hashlib
import re
import shutil
import struct
import subprocess
import tempfile
import threading
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from app.core.logger import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {
    "pdf",
    "ppt",
    "pptx",
    "md",
    "markdown",
    "doc",
    "docx",
    "png",
    "jpg",
    "jpeg",
}

_IMAGE_MIN_WIDTH = 120
_IMAGE_MIN_HEIGHT = 60
_MAX_IMAGES_PER_PAGE = 8
_HEADING_STYLE = re.compile(r"^(?:heading|标题)\s*(\d+)$", re.IGNORECASE)

_ZERO_WIDTH = re.compile(r"[​-‏‪-‮﻿]")
_MULTI_BLANK = re.compile(r"\n{3,}")
_SPACE_BEFORE_NEWLINE = re.compile(r"[ \t]+\n")
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_MEANINGFUL = re.compile(r"[一-鿿A-Za-z0-9]")

_ocr_engine: Any = None
_ocr_lock = threading.Lock()


class DocumentParseError(Exception):
    """文档解析失败（空文档、格式损坏等）。"""


@dataclass
class ParsedBlock:
    text: str
    title: str | None = None
    section: str | None = None
    page: int | None = None


@dataclass
class ParsedDocument:
    title: str
    blocks: list[ParsedBlock] = field(default_factory=list)

    @property
    def text_length(self) -> int:
        return sum(len(block.text) for block in self.blocks)


def clean_text(text: str) -> str:
    """统一换行、去除零宽字符与多余空白，保留段落结构。"""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _ZERO_WIDTH.sub("", text)
    text = _SPACE_BEFORE_NEWLINE.sub("\n", text)
    text = _MULTI_BLANK.sub("\n\n", text)
    lines = [line.strip() for line in text.split("\n")]
    return "\n".join(lines).strip()


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("文本编码无法识别，请使用 UTF-8 编码的文档")


def _paragraphs(text: str) -> list[str]:
    return [part.strip() for part in re.split(r"\n\s*\n", text) if part.strip()]


# --------------------------------------------------------------------------- #
# Markdown
# --------------------------------------------------------------------------- #
def parse_markdown(path: Path, title_hint: str | None = None) -> ParsedDocument:
    raw = _read_text_file(path)
    title = title_hint or path.stem
    title_locked = False
    current_section: str | None = None
    blocks: list[ParsedBlock] = []
    buffer: list[str] = []
    in_code_block = False

    def flush() -> None:
        if not buffer:
            return
        text = clean_text("\n".join(buffer))
        buffer.clear()
        if text:
            blocks.append(
                ParsedBlock(text=text, title=title, section=current_section)
            )

    for line in raw.split("\n"):
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            buffer.append(line.rstrip())
            continue

        heading = None if in_code_block else _HEADING.match(stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            heading_text = heading.group(2).strip()
            if level == 1 and not title_locked:
                title = heading_text
                title_locked = True
            else:
                current_section = heading_text
            continue

        if not stripped and not in_code_block:
            flush()
            continue

        if in_code_block:
            buffer.append(line.rstrip())
        else:
            buffer.append(stripped)

    flush()
    return ParsedDocument(title=title, blocks=blocks)


# --------------------------------------------------------------------------- #
# PDF（含扫描页 OCR）
# --------------------------------------------------------------------------- #
def _get_ocr_engine() -> Any:
    global _ocr_engine
    if _ocr_engine is None:
        with _ocr_lock:
            if _ocr_engine is None:
                from rapidocr_onnxruntime import RapidOCR

                logger.info("正在加载 RapidOCR 模型")
                _ocr_engine = RapidOCR()
    return _ocr_engine


def ocr_image(image: Any) -> list[str]:
    """对 OpenCV BGR 图像执行 OCR，返回按阅读顺序的文本行。"""
    result, _ = _get_ocr_engine()(image)
    if not result:
        return []
    lines: list[str] = []
    for item in result:
        if len(item) >= 2 and str(item[1]).strip():
            lines.append(str(item[1]).strip())
    return lines


def _decode_image(data: bytes) -> Any:
    import cv2
    import numpy as np

    array = np.frombuffer(data, dtype=np.uint8)
    image = cv2.imdecode(array, cv2.IMREAD_COLOR)
    if image is None:
        raise DocumentParseError("图片解码失败，文件可能已损坏")
    return image


def _ocr_image_bytes(data: bytes, seen: set[str] | None = None) -> str:
    """对文档内嵌图片做 OCR；跳过装饰图、重复图与过小图片。"""
    if not data:
        return ""
    digest = hashlib.sha1(data).hexdigest()
    if seen is not None:
        if digest in seen:
            return ""
        seen.add(digest)
    try:
        image = _decode_image(data)
    except Exception:
        logger.debug("内嵌图片解码失败，已跳过", exc_info=True)
        return ""
    height, width = image.shape[:2]
    if width < _IMAGE_MIN_WIDTH or height < _IMAGE_MIN_HEIGHT:
        return ""
    try:
        lines = ocr_image(image)
    except Exception:
        logger.warning("内嵌图片 OCR 失败，已跳过", exc_info=True)
        return ""
    return clean_text("\n".join(line for line in lines if line.strip()))


def ocr_image_bytes(data: bytes, seen: set[str] | None = None) -> str:
    """对图片字节执行 OCR，供 MinerU 等外部解析器复用。"""
    return _ocr_image_bytes(data, seen)


def _novel_ocr_lines(ocr_text: str, existing: str) -> list[str]:
    """去掉 OCR 与正文重复的行，避免同一句话重复进入向量。"""
    compact_existing = re.sub(r"\s+", "", existing)
    seen: set[str] = set()
    novel: list[str] = []
    for line in ocr_text.splitlines():
        line = line.strip()
        key = re.sub(r"\s+", "", line)
        if len(key) < 2 or key in seen or key in compact_existing:
            continue
        seen.add(key)
        novel.append(line)
    return novel


def _document_from_plain_text(
    text: str, title_hint: str | None = None, default_title: str = "文档"
) -> ParsedDocument:
    cleaned = clean_text(text)
    title = title_hint or default_title
    blocks = [
        ParsedBlock(text=paragraph, title=title)
        for paragraph in _paragraphs(cleaned)
    ]
    return ParsedDocument(title=title, blocks=blocks)


def parse_pdf(path: Path, title_hint: str | None = None) -> ParsedDocument:
    import pymupdf as fitz

    doc = fitz.open(str(path))
    try:
        meta_title = (doc.metadata or {}).get("title") or ""
        title = meta_title.strip() or title_hint or path.stem
        blocks: list[ParsedBlock] = []
        ocr_pages = 0
        ocr_images = 0
        seen_images: set[str] = set()

        for page_number, page in enumerate(doc, start=1):
            text = clean_text(page.get_text("text") or "")
            image_texts: list[str] = []
            if len(_MEANINGFUL.findall(text)) < 20:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                ocr_text = clean_text("\n".join(ocr_image(_decode_image(pixmap.tobytes("png")))))
                if ocr_text:
                    ocr_pages += 1
                    text = clean_text(f"{text}\n\n{ocr_text}") if text else ocr_text
            else:
                # 正文页中的尺码表、流程图、规则截图同样需要进入检索
                for image_info in page.get_images(full=True)[:_MAX_IMAGES_PER_PAGE]:
                    try:
                        extracted = doc.extract_image(image_info[0])
                    except Exception:
                        continue
                    image_ocr = _ocr_image_bytes(
                        extracted.get("image") or b"", seen_images
                    )
                    novel = _novel_ocr_lines(image_ocr, text)
                    if novel:
                        ocr_images += 1
                        image_texts.append("图片文字：\n" + "\n".join(novel))

            for paragraph in _paragraphs(text):
                blocks.append(
                    ParsedBlock(text=paragraph, title=title, page=page_number)
                )
            for image_text in image_texts:
                blocks.append(
                    ParsedBlock(text=image_text, title=title, page=page_number)
                )

        if not blocks:
            raise DocumentParseError("PDF 中未提取到有效文本（可能为空白或加密文档）")

        # 元数据缺失且未指定文件名提示时，用首页较短的首行作为标题
        if not meta_title.strip() and not title_hint:
            first_line = blocks[0].text.split("\n", 1)[0].strip()
            if 2 <= len(first_line) <= 40:
                title = first_line
                for block in blocks:
                    block.title = title

        logger.info(
            "PDF 解析完成 file=%s pages=%d blocks=%d ocr_pages=%d ocr_images=%d",
            path.name,
            doc.page_count,
            len(blocks),
            ocr_pages,
            ocr_images,
        )
        return ParsedDocument(title=title, blocks=blocks)
    finally:
        doc.close()


# --------------------------------------------------------------------------- #
# PPT / PPTX
# --------------------------------------------------------------------------- #
def _iter_shape_pictures(shapes: Any):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shape_pictures(shape.shapes)
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            yield shape


def _extract_pptx(path: Path, title_hint: str | None = None) -> ParsedDocument:
    from pptx import Presentation

    presentation = Presentation(str(path))
    title = title_hint or path.stem
    blocks: list[ParsedBlock] = []
    seen_images: set[str] = set()

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            slide_title = clean_text(slide.shapes.title.text)
            if slide_number == 1 and slide_title:
                title = slide_title

        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    content = clean_text(paragraph.text)
                    if content:
                        texts.append(content)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [clean_text(cell.text).replace("\n", " ") for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        texts.append(line)

        # 幻灯片中的商品卖点图、活动规则图、尺码示意图片
        for picture in _iter_shape_pictures(slide.shapes):
            try:
                image_ocr = _ocr_image_bytes(picture.image.blob, seen_images)
            except Exception:
                continue
            novel = _novel_ocr_lines(image_ocr, "\n".join(texts))
            if novel:
                texts.append("图片文字：\n" + "\n".join(novel))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = clean_text(slide.notes_slide.notes_text_frame.text)
            if notes:
                texts.append(f"备注：{notes}")

        # 标题已在 section 中体现，避免正文重复
        if slide_title:
            texts = [text for text in texts if text != slide_title]

        body = clean_text("\n".join(texts))
        if body:
            blocks.append(
                ParsedBlock(
                    text=body,
                    title=title,
                    section=slide_title or None,
                    page=slide_number,
                )
            )

    return ParsedDocument(title=title, blocks=blocks)


def _extract_legacy_ppt(path: Path, title_hint: str | None = None) -> ParsedDocument:
    """旧版 .ppt 采用 OLE 二进制格式，尽力提取 TextCharsAtom / TextBytesAtom。"""
    import olefile

    if not olefile.isOleFile(str(path)):
        raise DocumentParseError("PPT 文件格式无法识别")
    try:
        ole = olefile.OleFileIO(str(path))
        stream = ole.openstream("PowerPoint Document")
        data = stream.read()
    except Exception as exc:  # noqa: BLE001
        raise DocumentParseError(f"PPT 文件解析失败：{exc}") from exc

    texts: list[str] = []

    def walk(start: int, end: int, depth: int) -> None:
        if depth > 12:
            return
        offset = start
        while offset + 8 <= end:
            ver_instance, record_type, length = struct.unpack_from("<HHI", data, offset)
            body_start = offset + 8
            body_end = body_start + length
            if length < 0 or body_end > end:
                break
            version = ver_instance & 0x000F
            if record_type == 0x0FA0:  # TextCharsAtom (UTF-16LE)
                text = data[body_start:body_end].decode("utf-16-le", errors="ignore")
                if text.strip():
                    texts.append(text)
            elif record_type == 0x0FA8:  # TextBytesAtom (ANSI)
                text = data[body_start:body_end].decode("latin-1", errors="ignore")
                if text.strip():
                    texts.append(text)
            elif version == 0x0F:
                walk(body_start, body_end, depth + 1)
            offset = body_end

    walk(0, len(data), 0)
    cleaned = clean_text("\n".join(texts))
    title = title_hint or path.stem
    blocks = [ParsedBlock(text=paragraph, title=title) for paragraph in _paragraphs(cleaned)]
    return ParsedDocument(title=title, blocks=blocks)


def parse_ppt(path: Path, file_type: str, title_hint: str | None = None) -> ParsedDocument:
    if file_type == "pptx":
        return _extract_pptx(path, title_hint)
    return _extract_legacy_ppt(path, title_hint)


# --------------------------------------------------------------------------- #
# 图片
# --------------------------------------------------------------------------- #
def parse_image(path: Path, title_hint: str | None = None) -> ParsedDocument:
    image = _decode_image(path.read_bytes())
    lines = ocr_image(image)
    text = clean_text("\n".join(lines))
    if not text:
        raise DocumentParseError("图片中未识别到文字，请上传清晰度更高的图片")
    title = title_hint or path.stem
    blocks = [ParsedBlock(text=paragraph, title=title) for paragraph in _paragraphs(text)]
    return ParsedDocument(title=title, blocks=blocks)


# --------------------------------------------------------------------------- #
# Word (.docx / .doc)
# --------------------------------------------------------------------------- #
def _docx_heading_level(paragraph: Any) -> int:
    style = getattr(paragraph, "style", None)
    name = (getattr(style, "name", "") or "").strip()
    if not name:
        return 0
    match = _HEADING_STYLE.match(name)
    if match:
        return int(match.group(1))
    if name.lower() in {"title", "标题"}:
        return 1
    return 0


def _docx_element_images(
    element: Any, rels: Any, existing: str, seen: set[str]
) -> list[str]:
    from docx.oxml.ns import qn

    texts: list[str] = []
    for blip in element.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        if not rid:
            continue
        part = rels.get(rid)
        if part is None:
            continue
        image_ocr = _ocr_image_bytes(getattr(part, "blob", b""), seen)
        novel = _novel_ocr_lines(image_ocr, existing)
        if novel:
            texts.append("图片文字：\n" + "\n".join(novel))
    return texts


def _docx_table_text(table: Any) -> str:
    lines: list[str] = []
    for row in table.rows:
        cells = [clean_text(cell.text).replace("\n", " ") for cell in row.cells]
        line = " | ".join(cell for cell in cells if cell)
        if line.replace("|", "").strip():
            lines.append(line)
    return "\n".join(lines)


def _extract_docx(path: Path, title_hint: str | None = None) -> ParsedDocument:
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = DocxDocument(str(path))
    title = title_hint or path.stem
    title_locked = False
    section: str | None = None
    blocks: list[ParsedBlock] = []
    seen_images: set[str] = set()

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = clean_text(paragraph.text)
            level = _docx_heading_level(paragraph)
            images = _docx_element_images(
                child, document.part.related_parts, text, seen_images
            )
            if level:
                if level == 1 and text and not title_locked:
                    title = text
                    title_locked = True
                    section = None
                elif level >= 2 and text:
                    section = text
                if text:
                    blocks.append(
                        ParsedBlock(text=text, title=title, section=section)
                    )
                for image_text in images:
                    blocks.append(
                        ParsedBlock(text=image_text, title=title, section=section)
                    )
                continue
            body = text
            if images:
                body = clean_text("\n".join([body, *images])) if body else "\n".join(images)
            if body:
                blocks.append(ParsedBlock(text=body, title=title, section=section))
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            table_text = clean_text(_docx_table_text(table))
            image_texts = _docx_element_images(
                child, document.part.related_parts, table_text, seen_images
            )
            combined = table_text
            if image_texts:
                combined = clean_text("\n".join([combined, *image_texts]))
            if combined:
                blocks.append(
                    ParsedBlock(text=combined, title=title, section=section)
                )

    return ParsedDocument(title=title, blocks=blocks)


def _find_soffice() -> str | None:
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    return next((path for path in candidates if path and Path(path).exists()), None)


def _find_antiword() -> str | None:
    candidates = [
        shutil.which("antiword"),
        r"C:\Program Files\Git\mingw64\bin\antiword.exe",
        r"C:\Program Files (x86)\GnuWin32\bin\antiword.exe",
        "/usr/bin/antiword",
        "/usr/local/bin/antiword",
    ]
    return next((path for path in candidates if path and Path(path).exists()), None)


def _docx_from_soffice(
    path: Path, title_hint: str | None
) -> ParsedDocument | None:
    soffice = _find_soffice()
    if not soffice:
        return None
    with tempfile.TemporaryDirectory(prefix="atf_doc_") as tmp:
        try:
            result = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    tmp,
                    str(path),
                ],
                capture_output=True,
                text=True,
                timeout=90,
                check=False,
            )
        except Exception:
            logger.warning("LibreOffice 转换 .doc 失败，尝试其他方式", exc_info=True)
            return None
        converted = Path(tmp) / f"{path.stem}.docx"
        if result.returncode != 0 or not converted.exists():
            logger.warning(
                "LibreOffice 未生成 docx file=%s code=%s stderr=%s",
                path.name,
                result.returncode,
                (result.stderr or "")[:300],
            )
            return None
        return _extract_docx(converted, title_hint)


def _antiword_text(path: Path) -> str:
    antiword = _find_antiword()
    if not antiword:
        return ""
    # -w 0 关闭 antiword 自动换行，避免中英文混排时把表格单元格拆断
    for extra in (["-w", "0", "-m", "UTF-8.txt"], ["-w", "0"], ["-m", "UTF-8.txt"], []):
        try:
            result = subprocess.run(
                [antiword, *extra, str(path)],
                capture_output=True,
                timeout=60,
                check=False,
            )
        except Exception:
            logger.debug("antiword 调用失败", exc_info=True)
            continue
        if result.returncode == 0:
            text = result.stdout.decode("utf-8", errors="ignore")
            if len(_MEANINGFUL.findall(text)) >= 20:
                return text
    return ""


def _ole_doc_text(path: Path) -> str:
    """旧版 .doc 的最后兜底：从 WordDocument 流中提取可读文本。"""
    import olefile

    if not olefile.isOleFile(str(path)):
        raise DocumentParseError("旧版 Word 文件格式无法识别")
    ole = olefile.OleFileIO(str(path))
    try:
        if not ole.exists("WordDocument"):
            raise DocumentParseError("旧版 Word 文件缺少正文数据")
        data = ole.openstream("WordDocument").read()
    finally:
        ole.close()

    best = ""
    for encoding in ("utf-16-le", "gb18030"):
        decoded = data.decode(encoding, errors="ignore")
        chunks = re.findall(
            r"[\u4e00-\u9fffA-Za-z0-9，。！？；：、（）()《》【】%\-—\s]{8,}",
            decoded,
        )
        text = clean_text("\n".join(chunks))
        if len(_MEANINGFUL.findall(text)) > len(_MEANINGFUL.findall(best)):
            best = text
    return best


def _extract_legacy_doc(path: Path, title_hint: str | None = None) -> ParsedDocument:
    converted = _docx_from_soffice(path, title_hint)
    if converted is not None:
        logger.info("旧版 .doc 已通过 LibreOffice 转换解析 file=%s", path.name)
        return converted

    text = _antiword_text(path)
    if text:
        logger.info("旧版 .doc 已通过 antiword 解析 file=%s", path.name)
        return _document_from_plain_text(
            text.replace("\f", "\n\n"), title_hint, default_title=path.stem
        )

    text = _ole_doc_text(path)
    if text:
        logger.warning("旧版 .doc 使用 OLE 兜底解析，格式还原度有限 file=%s", path.name)
        return _document_from_plain_text(
            text, title_hint, default_title=path.stem
        )

    raise DocumentParseError(
        "旧版 .doc 解析失败，请安装 LibreOffice 或 antiword，"
        "也可以将文件另存为 .docx 后重新上传"
    )


# --------------------------------------------------------------------------- #
# 统一入口
# --------------------------------------------------------------------------- #
def parse_document(
    path: str | Path, file_type: str, title_hint: str | None = None
) -> ParsedDocument:
    path = Path(path)
    if not path.exists():
        raise DocumentParseError("文件不存在")
    extension = file_type.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError(f"不支持的文件类型：{extension}")

    if extension in ("md", "markdown"):
        parsed = parse_markdown(path, title_hint)
    elif extension == "pdf":
        parsed = parse_pdf(path, title_hint)
    elif extension in ("ppt", "pptx"):
        parsed = parse_ppt(path, extension, title_hint)
    elif extension == "docx":
        parsed = _extract_docx(path, title_hint)
    elif extension == "doc":
        parsed = _extract_legacy_doc(path, title_hint)
    else:
        parsed = parse_image(path, title_hint)

    parsed.blocks = [block for block in parsed.blocks if block.text.strip()]
    if not parsed.blocks:
        raise DocumentParseError(
            "未能从文档中提取到有效文本，请检查文件内容或格式后重试"
        )
    return parsed
