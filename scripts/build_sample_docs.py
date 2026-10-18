"""生成多格式测试样例：PDF / PPTX / 图片。

文档内容为逐条手工撰写，脚本仅负责格式转换，便于验证多格式解析链路。

用法：uv run python scripts/build_sample_docs.py
"""

from pathlib import Path

import pymupdf as fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "data" / "samples"

FONT_PATH = "C:/Windows/Fonts/msyh.ttc"
FONT_BOLD_PATH = "C:/Windows/Fonts/msyhbd.ttc"


def build_pdf() -> Path:
    lines = [
        "发票申请指南",
        "",
        "一、发票类型",
        "平台支持电子普通发票和增值税专用发票。",
        "电子普通发票与纸质发票具有同等法律效力。",
        "",
        "二、开票时间",
        "订单完成后三十天内可申请开票。",
        "电子普通发票在审核通过后三个工作日内开具。",
        "增值税专用发票在资料审核通过后五个工作日内开具。",
        "",
        "三、开票金额",
        "开票金额为订单实际支付金额。",
        "优惠券、积分抵扣和平台补贴不计入开票金额。",
        "",
        "四、发票换开",
        "发票开具后三十天内可申请换开。",
        "换开需先红冲原发票，通常需要三至五个工作日。",
        "",
        "五、常见问题",
        "问：退货后发票怎么处理？",
        "答：整单退货需配合红冲；部分退货按退款金额红冲后重新开具。",
    ]
    path = OUTPUT_DIR / "发票申请指南.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_font(fontname="msyh", fontfile=FONT_PATH)
    y = 64
    for line in lines:
        if line.startswith(("一、", "二、", "三、", "四、", "五、")):
            page.insert_text((64, y), line, fontname="msyh", fontsize=15, color=(0.2, 0.2, 0.6))
        else:
            page.insert_text((64, y), line, fontname="msyh", fontsize=12)
        y += 24
    document.save(path)
    document.close()
    return path


def build_pptx() -> Path:
    slides = [
        (
            "物流时效",
            [
                "现货商品付款后 48 小时内发出，大促期间延长至 72 小时。",
                "江浙沪皖 1 至 2 天送达，其他地区 2 至 6 天。",
                "新疆、西藏、内蒙古偏远地区 5 至 8 天。",
            ],
        ),
        (
            "运费标准",
            [
                "单笔订单实付满 199 元包邮。",
                "未满 199 元收取基础运费 10 元。",
                "偏远地区加收 15 元偏远运费。",
            ],
        ),
        (
            "物流异常处理",
            [
                "包裹 48 小时无更新，联系客服发起查询。",
                "确认丢件后可选补发或全额退款。",
                "外包装破损请当场拒收并拍照留证。",
            ],
        ),
    ]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title_layout = presentation.slide_layouts[0]
    content_layout = presentation.slide_layouts[1]

    first = presentation.slides.add_slide(title_layout)
    first.shapes.title.text = "物流说明"
    first.placeholders[1].text = "服装售前知识库 · 多格式解析样例"

    for title, bullets in slides:
        slide = presentation.slides.add_slide(content_layout)
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(20)
        for paragraph in frame.paragraphs:
            paragraph.font.size = Pt(20)
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(32)

    path = OUTPUT_DIR / "物流说明.pptx"
    presentation.save(path)
    return path


def build_image() -> Path:
    width, height = 900, 520
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(FONT_BOLD_PATH, 40)
    body_font = ImageFont.truetype(FONT_PATH, 26)
    draw.rectangle((0, 0, width, 90), fill=(79, 70, 229))
    draw.text((40, 22), "尺码速查", font=title_font, fill="white")
    lines = [
        "女装 M 码：身高 155-163cm，体重 48-55kg",
        "女装 L 码：身高 160-168cm，体重 55-62kg",
        "男装 L 码：身高 170-178cm，体重 65-75kg",
        "男装 XL 码：身高 175-182cm，体重 75-85kg",
        "修身版型建议选大一码，宽松版型可按正常码选择。",
        "尺码不合适支持七天无理由换货。",
    ]
    y = 130
    for line in lines:
        draw.text((40, y), line, font=body_font, fill=(30, 41, 59))
        y += 60
    path = OUTPUT_DIR / "尺码速查.png"
    image.save(path)
    return path


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for builder in (build_pdf, build_pptx, build_image):
        path = builder()
        print(f"已生成 {path.relative_to(PROJECT_ROOT)}")


if __name__ == "__main__":
    main()
